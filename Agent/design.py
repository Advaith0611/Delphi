import base64
import json
import logging
import mimetypes
import os
import shutil
from pathlib import Path
from typing import Any
from urllib.request import urlopen
from dotenv import load_dotenv
from agentspan.agents import Agent, EventType, shutdown, stream, tool


OPENROUTER_BASE_URL = "https://openrouter.ai/api/v1"
OPENROUTER_MODEL = "nvidia/nemotron-3-ultra-550b-a55b:free"
MODEL = f"openai/{OPENROUTER_MODEL}"

logger = logging.getLogger("design_agent")
ARTIFACT_ROOT = Path(__file__).with_name("artifacts")


def _artifact_path(filename: str) -> Path:
    """Return a safe path beneath the local artifacts directory."""
    path = Path(filename)
    if not path.name or path.name in {".", ".."} or path.name != filename:
        raise ValueError("filename must be a simple file name without directories")
    ARTIFACT_ROOT.mkdir(parents=True, exist_ok=True)
    return ARTIFACT_ROOT / path.name


def _result(path: Path, url: str | None = None) -> str:
    result = {"created": True, "path": str(path), "format": path.suffix.lower()}
    if url:
        result["url"] = url
        result["url_note"] = "Provider URLs may expire; use path as the durable local fallback."
    return json.dumps(result)


def _json_value(value: Any) -> Any:
    """Accept either a JSON string or an object already decoded by Agentspan."""
    if isinstance(value, str):
        return json.loads(value)
    return value


@tool
def create_text_artifact(filename: str, content: str) -> str:
    """
    Create a text-based artifact such as TXT, Markdown, HTML, CSS, JavaScript, or SVG.

    Use this for source files and vector artwork. The filename extension should
    match the content format. This tool creates or overwrites only a file inside
    the local artifacts directory.
    """
    path = _artifact_path(filename)
    path.write_text(content, encoding="utf-8")
    return _result(path)


@tool
def create_json_artifact(filename: str, content_json: Any) -> str:
    """
    Create a formatted JSON artifact from a JSON string.

    Use this for structured data or configuration files. The input must be valid
    JSON; the tool rejects invalid JSON instead of writing a malformed file.
    """
    path = _artifact_path(filename)
    data = _json_value(content_json)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return _result(path)


@tool
def create_spreadsheet(filename: str, title: str, rows_json: Any) -> str:
    """
    Create an XLSX spreadsheet with a title row and tabular data.

    rows_json must be a JSON array of arrays, or an array of objects. Use this
    for Excel-compatible workbooks with a readable header and basic styling.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    path = _artifact_path(filename)
    rows = _json_value(rows_json)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append([title])
    sheet["A1"].font = Font(bold=True, size=16)
    if rows and isinstance(rows[0], dict):
        headers = list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
    else:
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    return _result(path)


@tool
def create_document(filename: str, title: str, sections_json: Any) -> str:
    """
    Create a DOCX document with a title and ordered sections.

    sections_json must be a JSON array of objects containing a heading and body.
    """
    from docx import Document

    path = _artifact_path(filename)
    sections = _json_value(sections_json)
    document = Document()
    document.add_heading(title, level=0)
    for section in sections:
        document.add_heading(section["heading"], level=1)
        document.add_paragraph(section["body"])
    document.save(path)
    return _result(path)


@tool
def create_presentation(filename: str, title: str, slides_json: Any) -> str:
    """
    Create a PPTX presentation with a title slide and content slides.

    slides_json must be a JSON array of objects containing title and bullets,
    where bullets is an array of short strings.
    """
    from pptx import Presentation

    path = _artifact_path(filename)
    slides = _json_value(slides_json)
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    title_slide.placeholders[1].text = "Created by Pythios Design Agent"
    for slide_data in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = slide_data["title"]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, bullet in enumerate(slide_data.get("bullets", [])):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
    presentation.save(path)
    return _result(path)


@tool
def create_pdf(filename: str, title: str, sections_json: Any) -> str:
    """
    Create a PDF document with a title and ordered text sections.

    sections_json must be a JSON array of objects containing a heading and body.
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    path = _artifact_path(filename)
    sections = _json_value(sections_json)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 18)]
    for section in sections:
        story.extend([
            Paragraph(section["heading"], styles["Heading2"]),
            Paragraph(section["body"], styles["BodyText"]),
            Spacer(1, 10),
        ])
    SimpleDocTemplate(str(path), pagesize=letter).build(story)
    return _result(path)


@tool
def generate_image(prompt: str, filename: str = "generated.png"):
    """
    Generate an AI image and return the Image component tuple.

    The returned value is ``(image_data, seed)`` where ``image_data`` contains
    path, public URL or base64 data, byte size, original filename, MIME type,
    stream status, and metadata. The image is always copied into ``artifacts``
    when Gradio returns a local path or downloadable URL.
    """
    from gradio_client import Client

    client = Client(
        "black-forest-labs/FLUX.1-schnell",
        token=os.environ.get("HF_TOKEN") or os.environ.get("HUGGINGFACEHUB_API_TOKEN"),
    )
    result = client.predict(
        prompt=prompt,
        seed=0,
        randomize_seed=True,
        width=1024,
        height=1024,
        num_inference_steps=4,
        api_name="/infer",
    )

    if not isinstance(result, tuple) or len(result) != 2:
        raise RuntimeError(f"Unexpected image output from Gradio: {type(result).__name__}")

    raw_image, raw_seed = result
    if isinstance(raw_image, dict):
        image_metadata = raw_image
        source_path = image_metadata.get("path")
        source_url = image_metadata.get("url")
    elif isinstance(raw_image, (str, Path)):
        image_metadata = {}
        raw_image_value = str(raw_image)
        if raw_image_value.startswith(("http://", "https://", "data:")):
            source_path = None
            source_url = raw_image_value
        else:
            source_path = raw_image_value
            source_url = None
    else:
        image_metadata = {}
        source_path = getattr(raw_image, "path", None)
        source_url = getattr(raw_image, "url", None)
        if not source_path and not source_url:
            raise RuntimeError(
                "Gradio returned an unsupported image payload: "
                f"{type(raw_image).__name__}"
            )
    artifact_path = _artifact_path(filename)

    if source_path and Path(source_path).is_file():
        shutil.copyfile(source_path, artifact_path)
    elif isinstance(source_url, str) and source_url.startswith("data:"):
        _, encoded = source_url.split(",", 1)
        artifact_path.write_bytes(base64.b64decode(encoded))
    elif isinstance(source_url, str) and source_url.startswith(("http://", "https://")):
        with urlopen(source_url, timeout=120) as response:
            artifact_path.write_bytes(response.read())
    else:
        raise RuntimeError("Gradio returned no local image path or downloadable image URL")

    mime_type = image_metadata.get("mime_type") or mimetypes.guess_type(artifact_path.name)[0]
    image_data = {
        "path": str(artifact_path),
        "url": source_url if isinstance(source_url, str) else None,
        "size": artifact_path.stat().st_size,
        "orig_name": image_metadata.get("orig_name") or artifact_path.name,
        "mime_type": mime_type or "application/octet-stream",
        "is_stream": False,
        "meta": {
            **(image_metadata.get("meta") or {}),
            "prompt": prompt,
            "model": "black-forest-labs/FLUX.1-schnell",
        },
    }

    try:
        seed = float(raw_seed)
    except (TypeError, ValueError):
        seed = 0.0
    return image_data, seed



@tool
def generate_video(prompt: str, filename: str = "generated.mp4") -> str:
    """
    Generate an anchor image, animate it with LTX-2.5, and save an MP4 artifact.

    The image is generated first so the video has a stable composition and
    subject. LTX returns a four-element tuple; element zero is the temporary
    video filepath, which is copied into the local artifacts directory.
    """
    from gradio_client import Client, handle_file

    image_data, image_seed = generate_image(
        prompt=prompt,
        filename=f"{Path(filename).stem}_anchor.png",
    )
    image_path = image_data.get("path")
    if not image_path or not Path(image_path).is_file():
        raise RuntimeError("The anchor image was not saved locally")

    client = Client(
        "Lightricks/LTX-2.5",
        token=os.environ.get("HF_TOKEN") or os.environ.get("HUGGINGFACEHUB_API_TOKEN"),
    )
    result = client.predict(
        prompt=prompt,
        image_path=handle_file(image_path),
        height=1472,
        width=832,
        duration_s=2.0,
        seed=int(image_seed),
        decoder="conv",
        auto_len=False,
        randomize_seed=True,
        do_enhance=True,
        api_name="/run",
    )
    if not isinstance(result, tuple) or len(result) != 4:
        raise RuntimeError(f"Unexpected LTX-2.5 output: {result!r}")

    source_video, enhanced_prompt, video_seed, length = result
    if not source_video or not Path(source_video).is_file():
        raise RuntimeError("LTX-2.5 returned no local video filepath")

    artifact_path = _artifact_path(filename)
    shutil.copyfile(source_video, artifact_path)
    return json.dumps({
        "created": True,
        "path": str(artifact_path),
        "format": ".mp4",
        "anchor_image": image_data,
        "prompt": enhanced_prompt,
        "seed": video_seed,
        "length": length,
    })




def configure_logging() -> None:
    logging.basicConfig(
        level=os.environ.get("LOG_LEVEL", "INFO"),
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    )


def load_credentials() -> None:
    load_dotenv(Path(__file__).with_name(".env"))
    key = os.environ.get("OPENROUTER_API_KEY") or os.environ.get("OPENAI_API_KEY")
    if not key:
        raise RuntimeError("Set OPENROUTER_API_KEY or OPENAI_API_KEY in .env")
    os.environ["OPENAI_API_KEY"] = key


design_agent = Agent(
    name="design_agent",
    model=MODEL,
    base_url=OPENROUTER_BASE_URL,
    instructions=("You are the Design Agent for Pythios. You receive a sub-task "
"description and the relevant canvas region, including spatial layout, colors, "
"node sizes, and groupings. "
""
"Your job is to either (a) produce a new diagram or mockup that extends or "
"clarifies what's on the canvas, or (b) produce a direct patch to existing canvas "
"elements — never both unless the sub-task explicitly calls for both. "
""
"Preserve the visual language already present on the canvas: if the user uses a "
"particular color for 'data' nodes and another for 'process' nodes, continue that "
"convention in anything you add rather than introducing a new scheme. Respect "
"existing spatial groupings — new elements belonging to an existing cluster should "
"be positioned near that cluster, not placed arbitrarily. "
""
"Do not produce a full new diagram when a small patch (move a node, add one "
"connection, relabel a cluster) is sufficient — over-generating canvas content is "
"as much a failure mode as the wrong text output would be. "
""
"Return {patch_type: new_diagram|edit_patch, elements: [...], "
"canvas_region_target, rationale: one sentence on why this is the minimal "
"sufficient change}. "
""
"When the sub-task requests a deliverable file, use the appropriate artifact "
"tool. Files are created in the artifacts directory. Use text or JSON tools for "
"source and data files; use the document, spreadsheet, presentation, PDF, image, "
"and video tools for their corresponding formats. Do not claim a file was "
"created unless the tool returns success. Return created file paths and a concise "
"description of each artifact."),
    tools=[
        create_text_artifact,
        create_json_artifact,
        create_spreadsheet,
        create_document,
        create_presentation,
        create_pdf,
        generate_image,
        generate_video,
    ],
    max_tokens=4096,
)


def on_event(event) -> None:
    if event.type == EventType.TOOL_CALL:
        logger.info("tool call: %s %s", event.tool_name, event.args or {})
    elif event.type == EventType.TOOL_RESULT:
        logger.info("tool result: %s -> %s", event.tool_name, event.result)
    if event.type == EventType.DONE:
        logger.info("done: %s", event.execution_id)
    elif event.type == EventType.ERROR:
        logger.error("error: %s", event.content)


def run_design_agent(design_prompt: str) -> str:
    load_credentials()

    final_output = None
    for event in stream(design_agent, design_prompt):
        on_event(event)
        if event.type == EventType.DONE:
            final_output = event.output
            if (
                isinstance(final_output, dict)
                and final_output.get("finishReason") == "ERROR"
            ):
                raise RuntimeError(f"Design agent failed: {final_output}")
            break
        elif event.type == EventType.ERROR:
            raise RuntimeError(f"Design agent error: {event.content}")

    if final_output is None:
        raise RuntimeError("Design agent did not complete")

    return final_output


def main() -> None:
    try:
        print(run_design_agent("""
Test only the image and video generation tools. Do not create, overwrite, or
inspect any Markdown, JSON, DOCX, PDF, XLSX, or PPTX artifacts.

1. Call generate_image once with a concise visual prompt for a teal-and-charcoal
   product artifact showcase, saving it as showcase_test.png.
2. Call generate_video once with a concise motion prompt for the same visual
   theme, saving it as showcase_test.mp4. The video tool should generate its
   anchor image internally.

Report the exact returned result from each tool, including any exception or
failure. Do not claim success unless the tool returned success and the file path
was provided. Return a compact Markdown checklist containing only these two
tests.
"""))
    finally:
        shutdown()


if __name__ == "__main__":
    configure_logging()
    main()
